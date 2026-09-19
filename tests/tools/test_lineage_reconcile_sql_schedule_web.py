from __future__ import annotations

import unittest
from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from shared.lineage.domain import LineageEdge
from shared.lineage.environment_scope import LineageEnvironmentScope
from shared.lineage.reconciliation import (
    SCHEDULE_ACTIVE_SNAPSHOT_NOT_FOUND,
    ActiveSnapshotNotFoundError,
    LineageReconciliationResult,
    ReconciliationStatus,
    ScheduleLineageSnapshot,
    SQLBusinessLineageSnapshot,
    TargetSummaryStatus,
    reconcile_lineage_snapshots,
)
from shared.lineage.reconciliation_suppression import (
    SUPPRESSION_CLASSIFIER_VERSION,
    DWSReconciliationSuppressionRow,
    ReconciliationSuppression,
    ReconciliationSuppressionReason,
)
from shared.lineage.schedule import ScheduleLineageEdge
from tools.lineage.reconcile_sql_schedule_web import (
    EXPORT_HEADERS,
    EXPORT_SHEET_TITLE,
    SHOW_SELF_REFERENCE_OPTION,
    SHOW_SUPPRESSED_OPTION,
    ReconciliationErrorView,
    ReconciliationRowKind,
    ReconciliationViewModel,
    TargetReconciliationOutcome,
    _render_rows_html,
    build_environment_options,
    build_excel_bytes,
    build_export_filename,
    build_export_rows,
    build_reconciliation_view_model,
    build_summary,
    classify_reconciliation_row,
    is_self_reference_row,
    map_reconciliation_error,
    parse_target_tables,
    reconcile_target,
    reconcile_targets,
    resolve_display_options,
    sort_rows,
    status_to_label,
)

ENVIRONMENT = "DEMO_DEV"
SQL_PROFILE = "DEMO_SQL_PROFILE"
SCHEDULE_PROFILE = "DEMO_SCHEDULE_PROFILE"
OBSERVED_AT = datetime(2026, 9, 10, 8, 9, 10, tzinfo=timezone.utc)


def make_scope() -> LineageEnvironmentScope:
    return LineageEnvironmentScope(
        name="demo_dev",
        environment=ENVIRONMENT,
        sql_source_profile=SQL_PROFILE,
        schedule_source_profile=SCHEDULE_PROFILE,
        label="示例开发环境",
        dws_profile="demo",
    )


def make_result(
    target: str = "DWM.RESULT",
    *,
    sql_sources: tuple[str, ...] = ("DWF.A", "DWF.B"),
    schedule_sources: tuple[str, ...] = ("DWF.A", "DWF.C"),
):
    sql_edges = tuple(
        LineageEdge(
            environment=ENVIRONMENT,
            source_profile=SQL_PROFILE,
            source_table=source,
            target_table=target,
            program_name=f"DEMO_SQL_{source.rsplit('.', 1)[-1]}",
            batch_id="batch-sql",
            observed_at=OBSERVED_AT,
        )
        for source in sql_sources
    )
    schedule_edges = tuple(
        ScheduleLineageEdge(
            environment=ENVIRONMENT,
            source_profile=SCHEDULE_PROFILE,
            process_name=f"DEMO_SCHEDULE_{source.rsplit('.', 1)[-1]}",
            project_version_key="DEMO_PROJECT:1.0",
            raw_source_table=source,
            raw_target_table=target,
            source_table=source,
            target_table=target,
        )
        for source in schedule_sources
    )
    return reconcile_lineage_snapshots(
        SQLBusinessLineageSnapshot(
            batch_id="batch-sql",
            edges=sql_edges,
            observed_at=OBSERVED_AT,
            snapshot_scope=((ENVIRONMENT, SQL_PROFILE),),
        ),
        ScheduleLineageSnapshot(
            batch_id="batch-schedule",
            edges=schedule_edges,
            observed_at=OBSERVED_AT,
        ),
        environment=ENVIRONMENT,
        sql_source_profile=SQL_PROFILE,
        schedule_source_profile=SCHEDULE_PROFILE,
        target_table=target,
    )


def make_empty_result(target: str = "DWM.RESULT"):
    return make_result(target, sql_sources=(), schedule_sources=())


def make_batch_result(targets: tuple[str, ...]):
    sql_edges = tuple(
        LineageEdge(
            environment=ENVIRONMENT,
            source_profile=SQL_PROFILE,
            source_table=source,
            target_table=target,
            program_name=f"DEMO_SQL_{source.rsplit('.', 1)[-1]}",
            batch_id="batch-sql",
            observed_at=OBSERVED_AT,
        )
        for target in targets
        for source in ("DWF.A", "DWF.B")
    )
    schedule_edges = tuple(
        ScheduleLineageEdge(
            environment=ENVIRONMENT,
            source_profile=SCHEDULE_PROFILE,
            process_name=f"DEMO_SCHEDULE_{source.rsplit('.', 1)[-1]}",
            project_version_key="DEMO_PROJECT:1.0",
            raw_source_table=source,
            raw_target_table=target,
            source_table=source,
            target_table=target,
        )
        for target in targets
        for source in ("DWF.A", "DWF.C")
    )
    return reconcile_lineage_snapshots(
        SQLBusinessLineageSnapshot(
            batch_id="batch-sql",
            edges=sql_edges,
            observed_at=OBSERVED_AT,
            snapshot_scope=((ENVIRONMENT, SQL_PROFILE),),
        ),
        ScheduleLineageSnapshot(
            batch_id="batch-schedule",
            edges=schedule_edges,
            observed_at=OBSERVED_AT,
        ),
        environment=ENVIRONMENT,
        sql_source_profile=SQL_PROFILE,
        schedule_source_profile=SCHEDULE_PROFILE,
        target_tables=targets,
    )


def make_evidence_result(target: str = "DM.A") -> LineageReconciliationResult:
    """Build a result mixing NORMAL rows, one suppressed source and one self edge."""

    return make_result(
        target,
        sql_sources=(target, "DWF.B", "DWF.A"),
        schedule_sources=("DWF.A", "DWF.C"),
    )


def make_suppression_row(
    result: LineageReconciliationResult,
    *,
    source_table: str = "DWF.B",
    target_table: str | None = None,
    environment: str | None = None,
    sql_source_profile: str | None = None,
    schedule_source_profile: str | None = None,
    sql_batch_id: str | None = None,
    schedule_batch_id: str | None = None,
    classifier_version: str = SUPPRESSION_CLASSIFIER_VERSION,
    is_active: bool = True,
) -> DWSReconciliationSuppressionRow:
    observed_at = result.sql_observed_at or OBSERVED_AT
    candidate = ReconciliationSuppression(
        environment=environment or result.environment,
        sql_source_profile=sql_source_profile or result.sql_source_profile,
        schedule_source_profile=schedule_source_profile
        or result.schedule_source_profile,
        source_table=source_table,
        target_table=target_table or result.target_summaries[0].target_table,
        raw_status=ReconciliationStatus.SQL_ONLY,
        suppression_reason=ReconciliationSuppressionReason.NO_INTERNAL_PROGRAM,
        sql_batch_id=sql_batch_id or result.sql_batch_id,
        schedule_batch_id=schedule_batch_id or result.schedule_batch_id,
        classifier_version=classifier_version,
        observed_at=observed_at,
    )
    return DWSReconciliationSuppressionRow(
        row_key=candidate.row_key,
        suppression_key=candidate.suppression_key,
        environment=candidate.environment,
        sql_source_profile=candidate.sql_source_profile,
        schedule_source_profile=candidate.schedule_source_profile,
        source_table=candidate.source_table,
        target_table=candidate.target_table,
        raw_status=candidate.raw_status,
        suppression_reason=candidate.suppression_reason,
        sql_batch_id=candidate.sql_batch_id,
        schedule_batch_id=candidate.schedule_batch_id,
        classifier_version=candidate.classifier_version,
        observed_at=observed_at,
        first_seen_at=observed_at,
        last_seen_at=observed_at,
        is_active=is_active,
        created_at=observed_at,
        updated_at=observed_at,
    )


class _SuppressionReader:
    def __init__(self, rows=(), error: Exception | None = None) -> None:
        self.rows = tuple(rows)
        self.error = error
        self.calls: list[dict[str, object]] = []
        self.publish_calls = 0

    def read_rows(self, **kwargs):
        self.calls.append(kwargs)
        if self.error is not None:
            raise self.error
        return self.rows

    def publish(self, *args, **kwargs):
        self.publish_calls += 1
        raise AssertionError("UI must not publish suppression rows")


def run_with_suppression(
    result: LineageReconciliationResult,
    rows=(),
    *,
    reader_error: Exception | None = None,
):
    reader = _SuppressionReader(rows, error=reader_error)
    outcomes = reconcile_targets(
        make_scope(),
        (result.target_summaries[0].target_table,),
        runner=lambda **kwargs: result,
        suppression_store=reader,
    )
    return outcomes[0], reader


def require_view_model(
    outcome: TargetReconciliationOutcome,
) -> ReconciliationViewModel:
    if outcome.view_model is None:
        raise AssertionError("expected a successful reconciliation outcome")
    return outcome.view_model


def visible_sources(view_model: ReconciliationViewModel) -> list[str]:
    return [row.source_table for row in view_model.visible_rows()]


class ReconcileSqlScheduleWebTests(unittest.TestCase):
    def test_environment_options_expose_only_enabled_environment_values(self):
        scope = make_scope()
        disabled = LineageEnvironmentScope(
            name="disabled",
            environment="DEMO_DISABLED",
            sql_source_profile="DEMO_SQL_DISABLED",
            schedule_source_profile="DEMO_SCHEDULE_DISABLED",
            label="停用环境",
            dws_profile="demo",
            enabled=False,
        )

        options = build_environment_options((disabled, scope))

        self.assertEqual(len(options), 1)
        self.assertEqual(
            options[0].as_pywebio_option(),
            {
                "label": "示例开发环境",
                "value": ENVIRONMENT,
            },
        )

    def test_status_labels_keep_domain_status_values(self):
        self.assertEqual(status_to_label(ReconciliationStatus.MATCH), "两边一致")
        self.assertEqual(
            status_to_label(ReconciliationStatus.SQL_ONLY),
            "SQL实际调用但调度未配置",
        )
        self.assertEqual(
            status_to_label(ReconciliationStatus.SCHEDULE_ONLY),
            "调度已配置但SQL未调用",
        )

    def test_view_model_sorts_differences_first_and_builds_summary(self):
        result = make_result()

        view_model = build_reconciliation_view_model(result)

        self.assertEqual(view_model.status, TargetSummaryStatus.DIFFERENT)
        self.assertEqual(view_model.status_label, "有差异")
        self.assertEqual(
            [row.kind for row in view_model.rows],
            [ReconciliationRowKind.NORMAL] * 3,
        )
        self.assertEqual(view_model.rows, view_model.visible_rows())
        self.assertEqual(
            [row.status for row in view_model.rows],
            [
                ReconciliationStatus.SQL_ONLY,
                ReconciliationStatus.SCHEDULE_ONLY,
                ReconciliationStatus.MATCH,
            ],
        )
        self.assertEqual(view_model.sql_source_profile, SQL_PROFILE)
        self.assertEqual(view_model.schedule_source_profile, SCHEDULE_PROFILE)
        self.assertEqual(view_model.sql_batch_id, "batch-sql")
        self.assertEqual(view_model.schedule_batch_id, "batch-schedule")
        self.assertEqual(
            view_model.summary,
            build_summary(result.rows),
        )
        self.assertEqual(view_model.summary.sql_actual_count, 2)
        self.assertEqual(view_model.summary.schedule_configured_count, 2)
        self.assertEqual(view_model.summary.match_count, 1)
        self.assertEqual(view_model.summary.sql_only_count, 1)
        self.assertEqual(view_model.summary.schedule_only_count, 1)

    def test_current_active_sql_only_suppression_is_hidden_from_view(self):
        result = make_result()
        suppression = make_suppression_row(result)

        outcome, reader = run_with_suppression(result, (suppression,))

        view_model = require_view_model(outcome)
        self.assertEqual(visible_sources(view_model), ["DWF.C", "DWF.A"])
        self.assertEqual(
            [(row.source_table, row.kind) for row in view_model.rows],
            [
                ("DWF.C", ReconciliationRowKind.NORMAL),
                ("DWF.A", ReconciliationRowKind.NORMAL),
                ("DWF.B", ReconciliationRowKind.SUPPRESSED),
            ],
        )
        suppressed_row = view_model.visible_rows(show_suppressed=True)[-1]
        self.assertEqual(suppressed_row.source_table, "DWF.B")
        self.assertTrue(suppressed_row.is_evidence)
        self.assertEqual(suppressed_row.status_label, "手工码值/静态来源（不参与对账）")
        self.assertEqual(suppressed_row.schedule_display, "—")
        self.assertEqual(suppressed_row.sql_display, "是")
        self.assertEqual(view_model.summary.sql_only_count, 0)
        self.assertEqual(view_model.status, TargetSummaryStatus.DIFFERENT)
        self.assertEqual(
            reader.calls,
            [
                {
                    "environment": ENVIRONMENT,
                    "sql_source_profile": SQL_PROFILE,
                    "schedule_source_profile": SCHEDULE_PROFILE,
                    "sql_batch_id": "batch-sql",
                    "schedule_batch_id": "batch-schedule",
                    "active_only": True,
                    "target_tables": ("DWM.RESULT",),
                    "suppression_reason": ReconciliationSuppressionReason.NO_INTERNAL_PROGRAM,
                    "classifier_version": SUPPRESSION_CLASSIFIER_VERSION,
                }
            ],
        )

    def test_match_is_not_hidden_by_a_suppression_identity(self):
        result = make_result()

        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.A", "DWM.RESULT")},
        )

        self.assertIn("DWF.A", visible_sources(view_model))
        matched_row = next(
            row for row in view_model.rows if row.source_table == "DWF.A"
        )
        self.assertIs(matched_row.kind, ReconciliationRowKind.NORMAL)
        self.assertEqual(view_model.summary.match_count, 1)

    def test_actionable_sql_only_is_not_hidden_by_another_identity(self):
        result = make_result()

        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.A", "DWM.RESULT")},
        )

        self.assertIn("DWF.B", visible_sources(view_model))
        actionable_row = next(
            row for row in view_model.rows if row.source_table == "DWF.B"
        )
        self.assertIs(actionable_row.kind, ReconciliationRowKind.NORMAL)
        self.assertEqual(view_model.summary.sql_only_count, 1)

    def test_schedule_only_is_not_hidden_by_a_suppression_identity(self):
        result = make_result()

        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.C", "DWM.RESULT")},
        )

        self.assertIn("DWF.C", visible_sources(view_model))
        schedule_row = next(
            row for row in view_model.rows if row.source_table == "DWF.C"
        )
        self.assertIs(schedule_row.kind, ReconciliationRowKind.NORMAL)
        self.assertEqual(view_model.summary.schedule_only_count, 1)

    def test_stale_sql_batch_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(result, sql_batch_id="batch-sql-old")

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_stale_schedule_batch_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(
            result,
            schedule_batch_id="batch-schedule-old",
        )

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_environment_mismatch_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(result, environment="DEMO_OTHER")

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_sql_profile_mismatch_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(
            result,
            sql_source_profile="DEMO_SQL_OTHER",
        )

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_schedule_profile_mismatch_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(
            result,
            schedule_source_profile="DEMO_SCHEDULE_OTHER",
        )

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_inactive_suppression_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(result, is_active=False)

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_unsupported_classifier_version_does_not_hide_sql_only(self):
        result = make_result()
        suppression = make_suppression_row(result)
        object.__setattr__(
            suppression,
            "classifier_version",
            "reconciliation-suppression-v0",
        )

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_suppression_reader_exception_keeps_raw_sql_only_visible(self):
        result = make_result()

        outcome, _ = run_with_suppression(
            result,
            reader_error=RuntimeError("audit table unavailable"),
        )

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_malformed_suppression_row_keeps_raw_sql_only_visible(self):
        result = make_result()

        outcome, _ = run_with_suppression(result, ("malformed-row",))

        self.assertIn("DWF.B", visible_sources(require_view_model(outcome)))

    def test_excel_excludes_current_suppression_and_keeps_other_statuses(self):
        result = make_result()
        suppression = make_suppression_row(result)
        outcome, _ = run_with_suppression(result, (suppression,))

        export_rows = build_export_rows((outcome,))

        self.assertNotIn("DWF.B", [row.source_table for row in export_rows])
        self.assertIn(
            (ReconciliationStatus.MATCH, "DWF.A"),
            [(row.status, row.source_table) for row in export_rows],
        )
        self.assertIn(
            (ReconciliationStatus.SCHEDULE_ONLY, "DWF.C"),
            [(row.status, row.source_table) for row in export_rows],
        )

    def test_excel_keeps_actionable_sql_only_when_match_identity_is_supplied(self):
        result = make_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.A", "DWM.RESULT")},
        )
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        export_rows = build_export_rows((outcome,))

        self.assertIn(
            (ReconciliationStatus.SQL_ONLY, "DWF.B"),
            [(row.status, row.source_table) for row in export_rows],
        )

    def test_raw_result_is_not_mutated_by_presentation_filter(self):
        result = make_result()
        original_rows = result.rows
        suppression = make_suppression_row(result)

        outcome, _ = run_with_suppression(result, (suppression,))

        self.assertEqual(result.rows, original_rows)
        self.assertEqual(result.sql_only_count, 1)
        view_model = require_view_model(outcome)
        self.assertNotIn("DWF.B", visible_sources(view_model))
        self.assertIn(
            "DWF.B",
            [row.source_table for row in view_model.rows],
        )

    def test_multi_target_suppression_is_independent_by_target(self):
        first = make_result("DWM.RESULT_A")
        reader = _SuppressionReader((make_suppression_row(first),))

        outcomes = reconcile_targets(
            make_scope(),
            ("DWM.RESULT_A", "DWM.RESULT_B"),
            runner=lambda **kwargs: make_batch_result(kwargs["target_tables"]),
            suppression_store=reader,
        )

        self.assertEqual(len(outcomes), 2)
        self.assertNotIn("DWF.B", visible_sources(require_view_model(outcomes[0])))
        self.assertIn("DWF.B", visible_sources(require_view_model(outcomes[1])))
        self.assertEqual(len(reader.calls), 1)
        self.assertEqual(
            reader.calls[0]["target_tables"],
            ("DWM.RESULT_A", "DWM.RESULT_B"),
        )

    def test_ui_suppression_reader_is_read_only(self):
        result = make_result()
        reader = _SuppressionReader((make_suppression_row(result),))

        outcomes = reconcile_targets(
            make_scope(),
            ("DWM.RESULT",),
            runner=lambda **kwargs: result,
            suppression_store=reader,
        )

        self.assertTrue(outcomes[0].succeeded)
        self.assertEqual(reader.publish_calls, 0)
        self.assertEqual(len(reader.calls), 1)

    def test_display_options_never_change_the_formal_summary(self):
        result = make_evidence_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )
        baseline_summary = view_model.summary
        baseline_status = view_model.status
        baseline_normal = [
            row.source_table
            for row in view_model.rows
            if row.kind is ReconciliationRowKind.NORMAL
        ]

        for show_suppressed in (False, True):
            for show_self_reference in (False, True):
                with self.subTest(
                    show_suppressed=show_suppressed,
                    show_self_reference=show_self_reference,
                ):
                    visible = view_model.visible_rows(
                        show_suppressed=show_suppressed,
                        show_self_reference=show_self_reference,
                    )
                    self.assertEqual(view_model.summary, baseline_summary)
                    self.assertEqual(view_model.status, baseline_status)
                    self.assertEqual(
                        [
                            row.source_table
                            for row in visible
                            if row.kind is ReconciliationRowKind.NORMAL
                        ],
                        baseline_normal,
                    )

        self.assertEqual(baseline_summary.match_count, 1)
        self.assertEqual(baseline_summary.sql_only_count, 0)
        self.assertEqual(baseline_summary.schedule_only_count, 1)
        self.assertEqual(baseline_status, TargetSummaryStatus.DIFFERENT)
        self.assertEqual(visible_sources(view_model), ["DWF.C", "DWF.A"])
        self.assertEqual(
            [
                row.source_table
                for row in view_model.visible_rows(
                    show_suppressed=True,
                    show_self_reference=True,
                )
            ],
            ["DWF.C", "DWF.A", "DWF.B", "DM.A"],
        )

    def test_only_evidence_rows_keep_target_consistent(self):
        result = make_result("DM.A", sql_sources=("DM.A", "DWF.B"), schedule_sources=())
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )

        self.assertEqual(result.sql_only_count, 2)
        self.assertEqual(result.target_summaries[0].sql_only_count, 2)
        self.assertEqual(view_model.visible_rows(), ())
        self.assertEqual(
            [row.kind for row in view_model.rows],
            [ReconciliationRowKind.SUPPRESSED, ReconciliationRowKind.SELF_REFERENCE],
        )
        self.assertEqual(view_model.summary.sql_only_count, 0)
        self.assertEqual(view_model.summary.match_count, 0)
        self.assertEqual(view_model.status, TargetSummaryStatus.CONSISTENT)
        self.assertEqual(view_model.status_label, "一致")
        self.assertEqual(
            len(
                view_model.visible_rows(
                    show_suppressed=True,
                    show_self_reference=True,
                )
            ),
            2,
        )

    def test_self_reference_is_hidden_by_default_and_expandable(self):
        result = make_result("DM.A", sql_sources=("DM.A",), schedule_sources=())

        view_model = build_reconciliation_view_model(result)

        self.assertEqual(view_model.self_reference_count, 1)
        self.assertEqual(view_model.suppressed_count, 0)
        self.assertEqual(view_model.visible_rows(), ())
        self.assertEqual(view_model.summary.sql_only_count, 0)
        self.assertEqual(view_model.status, TargetSummaryStatus.CONSISTENT)

        expanded = view_model.visible_rows(show_self_reference=True)
        self.assertEqual(len(expanded), 1)
        row = expanded[0]
        self.assertIs(row.kind, ReconciliationRowKind.SELF_REFERENCE)
        self.assertTrue(row.is_evidence)
        self.assertEqual(row.status_label, "自关联（不参与调度对账）")
        self.assertEqual(row.schedule_display, "—")
        self.assertEqual(row.sql_display, "是")

    def test_self_reference_uses_comparison_normalized_identity(self):
        result = make_result("DM.A", sql_sources=("DWS_DM.A",), schedule_sources=())

        row = result.rows[0]
        self.assertEqual(row.source_table, "DM.A")
        self.assertEqual(row.target_table, "DM.A")
        self.assertTrue(is_self_reference_row(row))
        self.assertIs(
            classify_reconciliation_row(row),
            ReconciliationRowKind.SELF_REFERENCE,
        )

        other = make_result("DM.B", sql_sources=("DWS_DM.A",), schedule_sources=())
        other_row = other.rows[0]
        self.assertEqual(
            (other_row.source_table, other_row.target_table),
            ("DM.A", "DM.B"),
        )
        self.assertFalse(is_self_reference_row(other_row))
        self.assertIs(
            classify_reconciliation_row(other_row),
            ReconciliationRowKind.NORMAL,
        )

    def test_self_reference_wins_over_suppression_and_is_shown_once(self):
        result = make_result("DM.A", sql_sources=("DM.A",), schedule_sources=())
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DM.A", "DM.A")},
        )

        self.assertEqual(len(view_model.rows), 1)
        self.assertIs(view_model.rows[0].kind, ReconciliationRowKind.SELF_REFERENCE)
        self.assertEqual(view_model.suppressed_count, 0)
        self.assertEqual(view_model.self_reference_count, 1)
        self.assertEqual(view_model.visible_rows(show_suppressed=True), ())
        self.assertEqual(
            len(
                view_model.visible_rows(
                    show_suppressed=True,
                    show_self_reference=True,
                )
            ),
            1,
        )

    def test_resolve_display_options_accepts_checkbox_values_and_fails_closed(self):
        self.assertEqual(resolve_display_options(None), (False, False))
        self.assertEqual(resolve_display_options([]), (False, False))
        self.assertEqual(resolve_display_options("unknown"), (False, False))
        self.assertEqual(
            resolve_display_options([SHOW_SUPPRESSED_OPTION]),
            (True, False),
        )
        self.assertEqual(
            resolve_display_options((SHOW_SELF_REFERENCE_OPTION,)),
            (False, True),
        )
        self.assertEqual(
            resolve_display_options(
                [SHOW_SUPPRESSED_OPTION, SHOW_SELF_REFERENCE_OPTION]
            ),
            (True, True),
        )

    def test_sort_rows_prioritizes_sql_only_schedule_only_then_match(self):
        rows = make_result().rows

        sorted_rows = sort_rows(rows)

        self.assertEqual(
            [row.status for row in sorted_rows],
            [
                ReconciliationStatus.SQL_ONLY,
                ReconciliationStatus.SCHEDULE_ONLY,
                ReconciliationStatus.MATCH,
            ],
        )

    def test_export_rows_keep_all_statuses_and_sort_each_status_deterministically(self):
        result = make_result(
            sql_sources=("DWF.Z", "DWF.A"),
            schedule_sources=("DWF.Y", "DWF.B"),
        )
        view_model = build_reconciliation_view_model(result)
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        export_rows = build_export_rows((outcome,))

        self.assertEqual(
            [(row.status, row.source_table) for row in export_rows],
            [
                (ReconciliationStatus.SQL_ONLY, "DWF.A"),
                (ReconciliationStatus.SQL_ONLY, "DWF.Z"),
                (ReconciliationStatus.SCHEDULE_ONLY, "DWF.B"),
                (ReconciliationStatus.SCHEDULE_ONLY, "DWF.Y"),
            ],
        )

    def test_normal_sql_only_row_keeps_the_diff_styling(self):
        view_model = build_reconciliation_view_model(make_result())

        html = _render_rows_html(view_model)

        self.assertIn('<tr class="lineage-reconciliation-diff">', html)
        self.assertIn("SQL实际调用但调度未配置", html)
        self.assertNotIn('<tr class="lineage-reconciliation-evidence">', html)

    def test_evidence_rows_use_neutral_styling_instead_of_diff_red(self):
        result = make_evidence_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )

        html = _render_rows_html(
            view_model,
            show_suppressed=True,
            show_self_reference=True,
        )

        self.assertEqual(html.count('<tr class="lineage-reconciliation-evidence">'), 2)
        self.assertIn("手工码值/静态来源（不参与对账）", html)
        self.assertIn("自关联（不参与调度对账）", html)
        self.assertNotIn("SQL实际调用但调度未配置", html)

    def test_hidden_evidence_hint_reports_only_hidden_kinds(self):
        result = make_evidence_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )

        self.assertIn(
            "已隐藏（不参与对账）：静态来源 1 条，自关联 1 条。",
            _render_rows_html(view_model),
        )
        self.assertIn(
            "已隐藏（不参与对账）：自关联 1 条。",
            _render_rows_html(view_model, show_suppressed=True),
        )
        self.assertNotIn(
            "已隐藏",
            _render_rows_html(
                view_model,
                show_suppressed=True,
                show_self_reference=True,
            ),
        )

    def test_all_match_view_and_export_keep_complete_rows(self):
        result = make_result(
            sql_sources=("DWF.A", "DWF.B"),
            schedule_sources=("DWF.A", "DWF.B"),
        )
        view_model = build_reconciliation_view_model(result)
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        self.assertEqual(len(view_model.rows), 2)
        self.assertTrue(
            all(row.status is ReconciliationStatus.MATCH for row in view_model.rows)
        )
        self.assertEqual(_render_rows_html(view_model).count("两边一致"), 2)
        self.assertEqual(len(build_export_rows((outcome,))), 2)

    def test_excel_contains_fixed_headers_all_statuses_and_chinese_values(self):
        result = make_result()
        view_model = build_reconciliation_view_model(result)
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        workbook = load_workbook(
            BytesIO(build_excel_bytes(build_export_rows((outcome,))))
        )
        sheet = workbook[EXPORT_SHEET_TITLE]
        values = list(sheet.values)
        workbook.close()

        self.assertEqual(sheet.title, EXPORT_SHEET_TITLE)
        self.assertEqual(values[0], EXPORT_HEADERS)
        self.assertEqual(
            values[1:],
            [
                (
                    "DWM.RESULT",
                    "DWF.B",
                    "是",
                    "否",
                    "SQL实际调用但调度未配置",
                ),
                (
                    "DWM.RESULT",
                    "DWF.C",
                    "否",
                    "是",
                    "调度已配置但SQL未调用",
                ),
                ("DWM.RESULT", "DWF.A", "是", "是", "两边一致"),
            ],
        )

    def test_excel_excludes_both_evidence_kinds_by_default(self):
        result = make_evidence_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        export_rows = build_export_rows((outcome,))

        self.assertEqual(
            [(row.source_table, row.status) for row in export_rows],
            [
                ("DWF.C", ReconciliationStatus.SCHEDULE_ONLY),
                ("DWF.A", ReconciliationStatus.MATCH),
            ],
        )
        self.assertTrue(
            all(row.kind is ReconciliationRowKind.NORMAL for row in export_rows)
        )

    def test_excel_exports_expanded_evidence_with_presentation_labels(self):
        result = make_evidence_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        export_rows = build_export_rows(
            (outcome,),
            show_suppressed=True,
            show_self_reference=True,
        )
        workbook = load_workbook(BytesIO(build_excel_bytes(export_rows)))
        values = list(workbook.active.values)
        workbook.close()

        self.assertEqual(
            [row.source_table for row in export_rows],
            ["DWF.C", "DWF.A", "DWF.B", "DM.A"],
        )
        self.assertIn(
            ("DM.A", "DWF.B", "是", "—", "手工码值/静态来源（不参与对账）"),
            values[1:],
        )
        self.assertIn(
            ("DM.A", "DM.A", "是", "—", "自关联（不参与调度对账）"),
            values[1:],
        )
        flattened = [value for row in values for value in row]
        self.assertNotIn("SQL实际调用但调度未配置", flattened)

    def test_excel_display_options_are_independent(self):
        result = make_evidence_result()
        view_model = build_reconciliation_view_model(
            result,
            suppressed_edge_keys={("DWF.B", "DM.A")},
        )
        outcome = TargetReconciliationOutcome(
            target_table=view_model.target_table,
            view_model=view_model,
        )

        suppressed_only = build_export_rows((outcome,), show_suppressed=True)
        self.assertEqual(
            [row.source_table for row in suppressed_only],
            ["DWF.C", "DWF.A", "DWF.B"],
        )
        self_reference_only = build_export_rows((outcome,), show_self_reference=True)
        self.assertEqual(
            [row.source_table for row in self_reference_only],
            ["DWF.C", "DWF.A", "DM.A"],
        )

    def test_multi_target_export_keeps_evidence_inside_its_own_target(self):
        first = make_evidence_result("DM.A")
        second = make_result(
            "DM.B",
            sql_sources=("DM.B", "DWF.X"),
            schedule_sources=("DWF.X",),
        )
        outcomes = (
            TargetReconciliationOutcome(
                target_table="DM.A",
                view_model=build_reconciliation_view_model(
                    first,
                    suppressed_edge_keys={("DWF.B", "DM.A")},
                ),
            ),
            TargetReconciliationOutcome(
                target_table="DM.B",
                view_model=build_reconciliation_view_model(
                    second,
                    suppressed_edge_keys={("DWF.X", "DM.B")},
                ),
            ),
        )

        export_rows = build_export_rows(
            outcomes,
            show_suppressed=True,
            show_self_reference=True,
        )

        self.assertEqual(
            [(row.target_table, row.source_table, row.kind) for row in export_rows],
            [
                ("DM.A", "DWF.C", ReconciliationRowKind.NORMAL),
                ("DM.A", "DWF.A", ReconciliationRowKind.NORMAL),
                ("DM.B", "DWF.X", ReconciliationRowKind.NORMAL),
                ("DM.A", "DWF.B", ReconciliationRowKind.SUPPRESSED),
                ("DM.A", "DM.A", ReconciliationRowKind.SELF_REFERENCE),
                ("DM.B", "DM.B", ReconciliationRowKind.SELF_REFERENCE),
            ],
        )

    def test_multi_target_export_uses_one_sheet_and_excludes_failed_target(self):
        first = build_reconciliation_view_model(make_result("DWM.RESULT_A"))
        second = build_reconciliation_view_model(
            make_result(
                "DWM.RESULT_B", sql_sources=("DWF.X",), schedule_sources=("DWF.X",)
            )
        )
        outcomes = (
            TargetReconciliationOutcome(
                target_table=first.target_table,
                view_model=first,
            ),
            TargetReconciliationOutcome(
                target_table="DWM.FAIL",
                error=ReconciliationErrorView(
                    error_code=SCHEDULE_ACTIVE_SNAPSHOT_NOT_FOUND,
                    message="失败",
                ),
            ),
            TargetReconciliationOutcome(
                target_table=second.target_table,
                view_model=second,
            ),
        )

        export_rows = build_export_rows(outcomes)
        workbook = load_workbook(BytesIO(build_excel_bytes(export_rows)))
        sheet = workbook[EXPORT_SHEET_TITLE]
        values = list(sheet.values)
        workbook.close()

        self.assertEqual(workbook.sheetnames, [EXPORT_SHEET_TITLE])
        target_values = [row[0] for row in values[1:]]
        self.assertCountEqual(
            target_values,
            ["DWM.RESULT_A", "DWM.RESULT_A", "DWM.RESULT_A", "DWM.RESULT_B"],
        )
        self.assertNotIn("DWM.FAIL", target_values)

    def test_no_rows_means_no_export_payload(self):
        empty_view = build_reconciliation_view_model(make_empty_result())
        outcomes = (
            TargetReconciliationOutcome(
                target_table=empty_view.target_table,
                view_model=empty_view,
            ),
            TargetReconciliationOutcome(
                target_table="DWM.FAIL",
                error=ReconciliationErrorView(
                    error_code=SCHEDULE_ACTIVE_SNAPSHOT_NOT_FOUND,
                    message="失败",
                ),
            ),
        )

        self.assertEqual(build_export_rows(outcomes), ())

    def test_export_filename_is_safe_and_omits_profiles_and_batches(self):
        generated_at = datetime(2026, 9, 11, 15, 5, 0)

        single = build_export_filename(
            "DEV214",
            ("DWM.M_JJQD_LIST",),
            generated_at=generated_at,
        )
        multi = build_export_filename(
            "DEV214",
            ("DWM.RESULT_A", "DWM.RESULT_B"),
            generated_at=generated_at,
        )
        unsafe = build_export_filename(
            "DEV:214",
            ("DWM/M?JJQD|LIST",),
            generated_at="2026/09/11 15:05:00",
        )

        self.assertEqual(
            single,
            "lineage_reconciliation_DEV214_DWM_M_JJQD_LIST_20260911_150500.xlsx",
        )
        self.assertEqual(
            multi,
            "lineage_reconciliation_DEV214_20260911_150500.xlsx",
        )
        self.assertEqual(
            unsafe,
            "lineage_reconciliation_DEV_214_DWM_M_JJQD_LIST_2026_09_11_15_05_00.xlsx",
        )
        self.assertFalse(any(character in unsafe for character in ':\\/*?"<>|'))

    def test_reconcile_target_passes_resolved_split_scope_to_formal_runner(self):
        captured = {}

        def runner(**kwargs):
            captured.update(kwargs)
            return make_result(kwargs["target_table"])

        reconcile_target(make_scope(), " DWM.RESULT ", runner=runner)

        self.assertEqual(
            captured,
            {
                "dws_profile": "demo",
                "environment": ENVIRONMENT,
                "sql_source_profile": SQL_PROFILE,
                "schedule_source_profile": SCHEDULE_PROFILE,
                "target_table": "DWM.RESULT",
            },
        )

    def test_request_connection_is_forwarded_to_default_runner(self):
        captured = {}
        connection = object()

        def runner(**kwargs):
            captured.update(kwargs)
            return make_result(kwargs["target_table"])

        with patch(
            "tools.lineage.reconcile_sql_schedule_web.run_reconciliation",
            runner,
        ):
            outcomes = reconcile_targets(
                make_scope(),
                ("DWM.RESULT",),
                connection=connection,
            )

        self.assertTrue(outcomes[0].succeeded)
        self.assertIs(captured["connection"], connection)
        self.assertIn("timing", captured)

    def test_multiple_targets_use_one_batch_runner_and_preserve_order(self):
        calls: list[dict[str, object]] = []

        def runner(**kwargs):
            calls.append(kwargs)
            return make_batch_result(kwargs["target_tables"])

        outcomes = reconcile_targets(
            make_scope(),
            (" DWM.RESULT_A ", "DWM.RESULT_A", "DWA.RESULT_B"),
            runner=runner,
        )

        self.assertEqual(len(calls), 1)
        self.assertEqual(
            calls[0]["target_tables"],
            ("DWM.RESULT_A", "DWA.RESULT_B"),
        )
        self.assertEqual(
            [outcome.target_table for outcome in outcomes],
            ["DWM.RESULT_A", "DWA.RESULT_B"],
        )
        self.assertTrue(all(outcome.succeeded for outcome in outcomes))

    def test_fail_closed_error_mapping_preserves_both_snapshot_codes(self):
        sql_error = map_reconciliation_error(
            ActiveSnapshotNotFoundError("SQL_ACTIVE_SNAPSHOT_NOT_FOUND")
        )
        schedule_error = map_reconciliation_error(
            ActiveSnapshotNotFoundError(SCHEDULE_ACTIVE_SNAPSHOT_NOT_FOUND)
        )

        self.assertEqual(sql_error.error_code, "SQL_ACTIVE_SNAPSHOT_NOT_FOUND")
        self.assertIn("SQL active snapshot", sql_error.message)
        self.assertEqual(
            schedule_error.error_code,
            SCHEDULE_ACTIVE_SNAPSHOT_NOT_FOUND,
        )
        self.assertIn("Schedule active snapshot", schedule_error.message)

    def test_parse_target_tables_keeps_one_target_per_nonempty_line(self):
        self.assertEqual(
            parse_target_tables("\n DWM.TABLE_A \n\tDWM.TABLE_B\n"),
            ("DWM.TABLE_A", "DWM.TABLE_B"),
        )


if __name__ == "__main__":
    unittest.main()
