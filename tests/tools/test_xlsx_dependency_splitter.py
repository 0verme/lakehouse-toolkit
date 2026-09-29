import unittest
import zipfile
from io import BytesIO

from openpyxl import Workbook, load_workbook

from shared.graph.dependency_splitter import split_rows_preserving_dependencies
from tools.misc.xlsx_dependency_splitter import (
    _default_column,
    build_download_artifact,
    build_part_filename,
    load_first_worksheet,
)


def workbook_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class XlsxDependencySplitterTests(unittest.TestCase):
    def test_default_columns_detect_c_and_ab_without_fixing_other_choices(self):
        columns = ["job_name", "C", "AB"]

        self.assertEqual(_default_column(columns, "c"), "C")
        self.assertEqual(_default_column(columns, "ab"), "AB")
        self.assertIsNone(_default_column(columns, "dependencies"))

    def test_loads_only_first_worksheet_and_keeps_ordered_values(self):
        workbook = Workbook()
        first = workbook.active
        first.title = "作业"
        first.append(["job_name", "dependencies", "value"])
        first.append(["A", None, 10])
        first.append(["B", "33:A", 20])
        workbook.create_sheet("忽略的第二页").append(["other"])

        table = load_first_worksheet(workbook_bytes(workbook))

        self.assertEqual(table.columns, ["job_name", "dependencies", "value"])
        self.assertEqual(table.header_values, ["job_name", "dependencies", "value"])
        self.assertEqual(table.worksheet_title, "作业")
        self.assertEqual(len(table.rows), 2)
        self.assertEqual(table.rows[1]["dependencies"], "33:A")
        self.assertEqual(table.rows[1]["value"], 20)

    def test_rejects_blank_and_duplicate_headers(self):
        workbook = Workbook()
        workbook.active.append(["job_name", "job_name"])

        with self.assertRaisesRegex(ValueError, "列名重复"):
            load_first_worksheet(workbook_bytes(workbook))

    def test_empty_trailing_header_cells_are_ignored_but_data_without_header_fails(self):
        workbook = Workbook()
        workbook.active.append(["job_name", None])
        workbook.active.append(["A", "unexpected"])

        with self.assertRaisesRegex(ValueError, "没有表头的数据"):
            load_first_worksheet(workbook_bytes(workbook))

    def test_single_chunk_is_downloaded_as_xlsx(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["job_name", "dependencies"])
        worksheet.append(["A", None])
        table = load_first_worksheet(workbook_bytes(workbook))
        result = split_rows_preserving_dependencies(
            table.columns,
            table.rows,
            id_column="job_name",
            dependency_column="dependencies",
        )

        artifact = build_download_artifact("folder/source.xlsx", table, result)

        self.assertEqual(artifact.filename, "source_part001.xlsx")
        output = load_workbook(BytesIO(artifact.content), data_only=False)
        self.assertEqual(list(output.active.values), [("job_name", "dependencies"), ("A", None)])
        output.close()

    def test_multiple_chunks_are_bundled_as_xlsx_files_in_zip(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["job_name", "dependencies", "value"])
        worksheet.append(["A", None, 1])
        worksheet.append(["B", "33:A", 2])
        worksheet.append(["C", "33:A", 3])
        worksheet.append(["D", "33:A", 4])
        table = load_first_worksheet(workbook_bytes(workbook))
        result = split_rows_preserving_dependencies(
            table.columns,
            table.rows,
            id_column="job_name",
            dependency_column="dependencies",
            max_rows_per_chunk=2,
        )

        artifact = build_download_artifact("source.xlsx", table, result)

        self.assertEqual(artifact.filename, "source_split.zip")
        with zipfile.ZipFile(BytesIO(artifact.content)) as archive:
            self.assertEqual(
                archive.namelist(),
                [build_part_filename("source.xlsx", index) for index in range(1, 4)],
            )
            for filename in archive.namelist():
                output = load_workbook(BytesIO(archive.read(filename)), data_only=False)
                values = list(output.active.values)
                self.assertEqual(values[0], ("job_name", "dependencies", "value"))
                names = {row[0] for row in values[1:]}
                self.assertIn("A", names)
                self.assertLessEqual(len(values) - 1, 2)
                output.close()


if __name__ == "__main__":
    unittest.main()
