from __future__ import annotations

# The root-path bootstrap intentionally precedes imports used by direct script execution.
# ruff: noqa: I001

import sys
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

ROOT_DIR = Path(__file__).resolve().parents[2]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from shared.graph.dependency_splitter import (
    DependencySplitResult,
    split_rows_preserving_dependencies,
)

MAX_ROWS_PER_CHUNK = 500


@dataclass(frozen=True)
class XlsxTable:
    """Data from the first worksheet, without formatting or workbook styling."""

    columns: list[str]
    header_values: list[Any]
    rows: list[dict[str, Any]]
    worksheet_title: str


@dataclass(frozen=True)
class DownloadArtifact:
    filename: str
    content: bytes


def _is_blank_header(value: Any) -> bool:
    return value is None or not str(value).strip()


def load_first_worksheet(content: bytes) -> XlsxTable:
    """Load the first worksheet's header and non-empty data rows from XLSX bytes."""

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        if not workbook.worksheets:
            raise ValueError("XLSX 文件中没有 worksheet")
        worksheet = workbook.worksheets[0]
        row_iterator = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = list(next(row_iterator))
        except StopIteration as exc:
            raise ValueError("第一个 worksheet 为空，找不到表头") from exc

        while raw_headers and _is_blank_header(raw_headers[-1]):
            raw_headers.pop()
        if not raw_headers:
            raise ValueError("第一个 worksheet 的首行没有有效表头")
        if any(_is_blank_header(header) for header in raw_headers):
            raise ValueError("表头中存在空列名，请先补齐表头")

        columns = [str(header).strip() for header in raw_headers]
        seen_columns: set[str] = set()
        for column in columns:
            if column in seen_columns:
                raise ValueError(f"表头列名重复：{column}")
            seen_columns.add(column)

        rows: list[dict[str, Any]] = []
        for worksheet_row, values in enumerate(row_iterator, start=2):
            values = list(values)
            if not any(value is not None for value in values):
                continue
            if any(value is not None for value in values[len(columns) :]):
                raise ValueError(f"worksheet 第 {worksheet_row} 行存在没有表头的数据")
            padded_values = values[: len(columns)] + [
                None
            ] * max(0, len(columns) - len(values))
            rows.append(dict(zip(columns, padded_values, strict=True)))

        return XlsxTable(
            columns=columns,
            header_values=raw_headers,
            rows=rows,
            worksheet_title=worksheet.title,
        )
    finally:
        workbook.close()


def build_xlsx_bytes(table: XlsxTable, rows: list[dict[str, Any]]) -> bytes:
    """Write tabular values to a basic XLSX workbook, without style guarantees."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = table.worksheet_title[:31] or "Sheet1"
    worksheet.append(table.header_values)
    for row in rows:
        worksheet.append([row.get(column) for column in table.columns])

    buffer = BytesIO()
    try:
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()


def _input_stem(input_filename: str) -> str:
    normalized = str(input_filename or "workbook.xlsx").replace("\\", "/")
    stem = Path(normalized).name.rsplit(".", 1)[0].strip()
    return stem or "workbook"


def build_part_filename(input_filename: str, part_number: int) -> str:
    return f"{_input_stem(input_filename)}_part{part_number:03d}.xlsx"


def build_download_artifact(
    input_filename: str,
    table: XlsxTable,
    result: DependencySplitResult,
) -> DownloadArtifact | None:
    """Return one XLSX directly or a ZIP containing all XLSX parts."""

    if not result.chunks:
        return None
    if len(result.chunks) == 1:
        return DownloadArtifact(
            filename=build_part_filename(input_filename, 1),
            content=build_xlsx_bytes(table, result.chunks[0]),
        )

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for part_number, rows in enumerate(result.chunks, start=1):
            archive.writestr(
                build_part_filename(input_filename, part_number),
                build_xlsx_bytes(table, rows),
            )
    return DownloadArtifact(
        filename=f"{_input_stem(input_filename)}_split.zip",
        content=buffer.getvalue(),
    )


def _default_column(columns: list[str], name: str) -> str | None:
    return next((column for column in columns if column.casefold() == name), None)


def main() -> None:
    pywebio_input = __import__("pywebio.input", fromlist=["file_upload"])
    pywebio_output = __import__("pywebio.output", fromlist=["put_file"])
    file_upload = pywebio_input.file_upload
    input_group = pywebio_input.input_group
    select = pywebio_input.select
    number = pywebio_input.input
    put_file = pywebio_output.put_file
    put_markdown = pywebio_output.put_markdown
    put_text = pywebio_output.put_text

    from shared.ui.pywebio_helper import safe_put_error

    put_markdown("## 依赖关系 Excel 自动拆分")
    put_markdown(
        "上传包含唯一标识和依赖关系的 `.xlsx`。工具读取第一个 worksheet；"
        "依赖字段沿用仓库现有格式，例如 `33:JOB_A|33:JOB_B`。"
    )
    uploaded = file_upload("请选择 XLSX 文件", accept=".xlsx", required=True)
    if not uploaded:
        return

    input_filename = uploaded.get("filename", "workbook.xlsx")
    if not str(input_filename).lower().endswith(".xlsx"):
        safe_put_error("当前仅支持 .xlsx 文件")
        return
    try:
        table = load_first_worksheet(uploaded["content"])
    except Exception as exc:  # noqa: BLE001 - upload errors should be shown in the UI
        safe_put_error(exc)
        return

    id_default = _default_column(table.columns, "c")
    dependency_default = _default_column(table.columns, "ab")
    id_options = [
        {
            "label": column,
            "value": column,
            **({"selected": True} if column == id_default else {}),
        }
        for column in table.columns
    ]
    dependency_options = [
        {
            "label": column,
            "value": column,
            **({"selected": True} if column == dependency_default else {}),
        }
        for column in table.columns
    ]
    form = input_group(
        "拆分参数",
        [
            select("唯一标识列", name="id_column", options=id_options),
            select("依赖关系列", name="dependency_column", options=dependency_options),
            number(
                "单文件最大数据行数",
                name="max_rows_per_chunk",
                type=pywebio_input.NUMBER,
                value=MAX_ROWS_PER_CHUNK,
                min=1,
            ),
        ],
    )

    try:
        max_rows = int(form["max_rows_per_chunk"])
        result = split_rows_preserving_dependencies(
            table.columns,
            table.rows,
            id_column=form["id_column"],
            dependency_column=form["dependency_column"],
            max_rows_per_chunk=max_rows,
        )
    except Exception as exc:  # noqa: BLE001 - validation errors should be shown in the UI
        safe_put_error(exc)
        return

    diagnostics = result.diagnostics
    put_markdown(
        f"原始数据行数：**{diagnostics.original_data_rows}**  \n"
        f"有效唯一标识：**{diagnostics.valid_id_count}**  \n"
        f"缺失依赖：**{diagnostics.missing_dependency_count}**  \n"
        f"单文件上限：**{max_rows}** 行  \n"
        f"最终拆分：**{len(result.chunks)}** 份"
    )
    if diagnostics.empty_id_count:
        put_markdown(
            f"> ⚠️ 有 **{diagnostics.empty_id_count}** 行唯一标识为空，已从依赖图和拆分文件中排除。"
        )
    if diagnostics.missing_dependency_count:
        missing_names = "、".join(diagnostics.missing_dependencies)
        put_markdown(
            f"> ⚠️ 缺失依赖共 **{diagnostics.missing_dependency_count}** 个唯一名称，"
            f"出现 **{diagnostics.missing_dependency_references}** 次；未伪造对应记录：{missing_names}"
        )
    if diagnostics.cycles:
        rendered_cycles = "；".join(" → ".join(cycle) for cycle in diagnostics.cycles)
        put_markdown(
            f"> ⚠️ 检测到循环依赖（最多展示 20 个）；拆分按可达闭包保留环内记录：{rendered_cycles}"
        )
    put_markdown(
        "为保证每个拆分文件内部依赖完整，部分上游记录可能在多个文件中重复出现。"
    )

    if not result.chunks:
        put_text("没有有效唯一标识行，未生成下载文件。")
        return

    part_summary = [
        f"{build_part_filename(input_filename, index)}    {len(rows)} 行"
        for index, rows in enumerate(result.chunks, start=1)
    ]
    put_text("\n".join(part_summary))
    artifact = build_download_artifact(input_filename, table, result)
    if artifact is not None:
        put_file(artifact.filename, artifact.content, "下载拆分结果")


if __name__ == "__main__":
    from shared.ui.pywebio_helper import start_pywebio_app

    start_pywebio_app("依赖关系 Excel 自动拆分", main)
